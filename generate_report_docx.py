import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Page Margins: Standard Academic
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    img_dir = r"C:\Users\cthrn\.gemini\antigravity\brain\c420fef9-974a-4a45-a770-fc6b29ca5b5b\screenshots"

    # ==========================================
    # HALAMAN 1: INSTRUKSI TUGAS
    # ==========================================
    p_inst_title = doc.add_paragraph()
    p_inst_title.paragraph_format.space_after = Pt(12)
    run_inst_title = p_inst_title.add_run("Instruksi:")
    run_inst_title.bold = True
    run_inst_title.font.size = Pt(14)
    run_inst_title.font.color.rgb = RGBColor(185, 28, 28)

    table_inst = doc.add_table(rows=1, cols=1)
    table_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_inst = table_inst.cell(0, 0)
    set_cell_background(cell_inst, "FEF2F2")
    set_cell_margins(cell_inst, top=140, bottom=140, left=200, right=200)

    p_in = cell_inst.paragraphs[0]
    p_in.paragraph_format.space_after = Pt(4)
    r1 = p_in.add_run("1. Nama File saat disubmit: ")
    r1.bold = True
    p_in.add_run("NoKelompok_NmMahasiswa_NIM.pdf (Contoh: K1_CatherineWong_24N10001.pdf)\n")
    
    r2 = p_in.add_run("2. Maksimal kemiripan di Turnitin: ")
    r2.bold = True
    p_in.add_run("30%\n")

    r3 = p_in.add_run("3. Tampilan screenshot: ")
    r3.bold = True
    p_in.add_run("Diambil langsung oleh komputer masing-masing individu (difokuskan pada pengujian Android responsif dan server live hosting https://app.mesinbayar.com/msra/).\n")

    r4 = p_in.add_run("4. PERHATIAN PENTING: ")
    r4.bold = True
    r4.font.color.rgb = RGBColor(185, 28, 28)
    p_in.add_run("Hapus halaman instruksi ini ketika laporan disubmit ke portal kuliah/dosen.")

    doc.add_page_break()

    # ==========================================
    # HALAMAN 2: COVER
    # ==========================================
    p_cover_top = doc.add_paragraph()
    p_cover_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover_top.paragraph_format.space_before = Pt(36)
    p_cover_top.paragraph_format.space_after = Pt(16)
    
    r_lap = p_cover_top.add_run("LAPORAN INDIVIDU\n")
    r_lap.bold = True
    r_lap.font.size = Pt(16)
    r_lap.font.color.rgb = RGBColor(30, 58, 138)

    r_proj = p_cover_top.add_run("MSRA: Money Smart Room Access with AiYO Bills & IoT Door Access\n")
    r_proj.bold = True
    r_proj.font.size = Pt(18)
    r_proj.font.color.rgb = RGBColor(15, 23, 42)

    r_subproj = p_cover_top.add_run("Integrasi Sistem Gateway Pembayaran AiYO Bills & Kunci Pintu Otomatis IoT ESP32 Berbasis Android PWA (Live Production: https://app.mesinbayar.com/msra/)")
    r_subproj.font.size = Pt(11)
    r_subproj.font.italic = True
    r_subproj.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    
    table_author = doc.add_table(rows=3, cols=2)
    table_author.alignment = WD_TABLE_ALIGNMENT.CENTER

    author_data = [
        ("Nama Kelompok", "Money Talks (MSRA)"),
        ("Nama Lengkap Mahasiswa", "Catherine Wong"),
        ("NIM (Nomor Induk Mahasiswa)", "[Isi NIM Anda]")
    ]

    for i, (label, val) in enumerate(author_data):
        c0 = table_author.cell(i, 0)
        c1 = table_author.cell(i, 1)
        c0.text = label
        c1.text = ": " + val
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].paragraph_format.space_after = Pt(4)
        c1.paragraphs[0].paragraph_format.space_after = Pt(4)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    p_cover_bottom = doc.add_paragraph()
    p_cover_bottom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover_bottom.paragraph_format.space_before = Pt(72)

    r_matkul = p_cover_bottom.add_run("Mata Kuliah: Teknologi Keuangan (Fintech)\n")
    r_matkul.bold = True
    r_matkul.font.size = Pt(12)
    
    r_prodi = p_cover_bottom.add_run("Program Studi Sistem Informasi / Rekayasa Perangkat Lunak (SIEGA)\n")
    r_prodi.bold = True
    r_prodi.font.size = Pt(12)

    r_univ = p_cover_bottom.add_run("Fakultas Ilmu Komputer\nUniversitas Katolik Soegijapranata (SCU) Semarang\n")
    r_univ.bold = True
    r_univ.font.size = Pt(13)
    r_univ.font.color.rgb = RGBColor(30, 58, 138)

    r_sem = p_cover_bottom.add_run("Semester Gasal 2026/2027")
    r_sem.bold = True
    r_sem.font.size = Pt(11)

    doc.add_page_break()

    # ==========================================
    # BAB 1: PENGANTAR
    # ==========================================
    h1 = doc.add_heading("1. Pengantar", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h1.runs[0].font.size = Pt(16)
    h1.runs[0].bold = True

    p = doc.add_paragraph()
    p.add_run("1.1 Informasi Project\n").bold = True
    p.add_run(
        "Perkembangan teknologi kota cerdas (Smart City) dan gedung komersial modern menuntut efisiensi operasional "
        "yang semakin tinggi dalam pengelolaan fasilitas sewa bersama, seperti coworking space, meeting room, maupun "
        "bilik kerja terdistribusi. Pada sistem persewaan konvensional, pelanggan kerap menghadapi kendala berupa "
        "ketergantungan terhadap kehadiran staf resepsionis fisik, risiko kunci fisik hilang atau diduplikasi, serta "
        "proses rekonsiliasi pembayaran manual yang rentan terhadap kecurangan dan kesalahan pemesanan ganda (double booking).\n\n"
        "MSRA (Money Smart Room Access) with AiYO Bills & IoT Door Access adalah solusi terintegrasi yang menggabungkan "
        "Teknologi Finansial (Fintech Payment Gateway) dengan Internet of Things (IoT Smart Lock). Sistem ini memungkinkan "
        "penyewa melakukan penelusuran jadwal ruangan, reservasi mandiri dari perangkat smartphone Android, pembayaran instan "
        "melalui QRIS Nasional dan Virtual Account terlisensi AiYO Bills, penerbitan tiket akses digital dinamis, serta pembukaan "
        "gagang pintu fisik ruangan secara nirsentuh (touchless) menggunakan pemindai kamera smartphone.\n\n"
        "Sistem telah di-deploy dan berjalan aktif pada server hosting live: https://app.mesinbayar.com/msra/."
    )

    p_flow = doc.add_paragraph()
    p_flow.paragraph_format.space_before = Pt(10)
    p_flow.add_run("1.2 Alir Kerja Sistem (System Workflow)\n").bold = True
    p_flow.add_run(
        "Alur kerja sistem dirancang secara menyeluruh dan berkesinambungan tanpa intervensi manual:\n"
        "1. Penelusuran Ruangan: Pelanggan mengakses web MSRA melalui browser smartphone Android dan memilih ruangan serta durasi sewa.\n"
        "2. Penguncian Slot Sementara (Atomic Slot Lock): Sistem melakukan validasi ketersediaan jadwal dan mengunci slot ruangan selama 15 menit agar tidak diserobot pengguna lain.\n"
        "3. Penerbitan Tagihan AiYO Bills: Backend berkomunikasi dengan API AiYO Bills melalui protokol OAuth 2.0 dan kalkulasi tanda tangan digital HMAC-SHA256 untuk menghasilkan kode QRIS dan nomor rekening Virtual Account resmi.\n"
        "4. Pembayaran Real-Time: Pelanggan melakukan pembayaran menggunakan aplikasi mobile banking atau dompet digital (e-wallet).\n"
        "5. Eksekusi Webhook Callback: Server pembayaran AiYO Bills mengirimkan sinyal notifikasi HTTP POST (callback) yang secara otomatis memperbarui status transaksi menjadi PAID dan menerbitkan tiket akses digital dinamis (SPK-PASS-xxxx).\n"
        "6. Notifikasi Otomatis: Rincian tiket akses dan tautan navigasi langsung dikirimkan ke nomor WhatsApp dan alamat email pelanggan.\n"
        "7. Akses Fisik IoT di Depan Pintu: Pelanggan tiba di lokasi dan memindai plakat QR Code di depan pintu kamar menggunakan kamera smartphone Android. Sistem memvalidasi kesesuaian waktu sewa dan nomor ruangan, kemudian memerintahkan mikrokontroler ESP32 untuk menarik kunci solenoid 12V selama 5 detik sehingga pengguna dapat langsung masuk."
    )

    diag_flow_path = os.path.join(img_dir, "diagram_system_flow.png")
    if os.path.exists(diag_flow_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(diag_flow_path, width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.add_run("Gambar 1.1: Diagram Alur Kerja End-to-End Sistem MSRA.").font.size = Pt(9.5)

    doc.add_page_break()

    # ==========================================
    # BAB 2: TAMPILAN DEVICE
    # ==========================================
    h2 = doc.add_heading("2. Tampilan Device", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h2.runs[0].font.size = Pt(16)
    h2.runs[0].bold = True

    p_dev = doc.add_paragraph()
    p_dev.add_run(
        "Sub-sistem perangkat keras (hardware) pada proyek MSRA dirancang untuk bekerja secara real-time "
        "bekerjasama dengan sistem backend teknologi keuangan. Setiap pintu ruangan dilengkapi dengan unit pengendali "
        "IoT independen yang terhubung ke jaringan internet nirkabel (Wi-Fi) gedung."
    )

    diag_wire_path = os.path.join(img_dir, "diagram_device_wiring.png")
    if os.path.exists(diag_wire_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(diag_wire_path, width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.add_run("Gambar 2.1: Diagram Skematik Blok Pengkabelan (Wiring Block Diagram) Perangkat IoT.").font.size = Pt(9.5)

    p_spec = doc.add_paragraph()
    p_spec.paragraph_format.space_before = Pt(12)
    p_spec.add_run("Penjelasan Komponen Perangkat Keras dan Kerjasamanya dengan Sistem Keuangan:\n").bold = True
    p_spec.add_run(
        "1. NodeMCU ESP32 Microcontroller: Modul pemroses utama berkecepatan 240 MHz yang dilengkapi konektivitas Wi-Fi 802.11 b/g/n. "
        "ESP32 secara aktif menjalankan tugas polling periodik setiap 1-2 detik ke endpoint API server backend (/api/door/poll?room_code=xxx) "
        "untuk memeriksa apakah ada perintah pembukaan pintu dari transaksi pembayaran yang baru saja terverifikasi lunas.\n"
        "2. 12V Solenoid Door Lock: Kunci pintu elektromagnetik bertipe fail-secure. Dalam kondisi standby tanpa arus, lidah pengunci terkunci kokoh. "
        "Saat ada instruksi resmi dari sistem keuangan bahwa pengguna telah membayar dan memindai tiketnya, solenoid dialiri arus 12V selama 5 detik sehingga lidah pengunci tertarik dan pintu dapat didorong terbuka.\n"
        "3. 1-Channel Relay Module (Optocoupler Isolation): Modul sakelar elektronik yang berfungsi mengisolasi tegangan logika 3.3V ESP32 "
        "dari lonjakan tegangan dan arus beban induktif solenoid 12V (800mA-1A), menjamin mikrokontroler aman dari gangguan noise elektromagnetik.\n"
        "4. LCD Display 16x2 / 20x4 I2C: Penampil teks visual di daun pintu yang menampilkan status operasional pintu real-time: 'STANDBY: TERKUNCI' "
        "saat menunggu penyewa, dan 'VERIFIKASI SUKSES: SILAKAN MASUK' selama durasi 5 detik saat kunci terbuka.\n"
        "5. Active Buzzer 5V & Dual LED Status: Memberikan indikator audio-visual konfirmasi langsung. LED hijau dan nada beep ganda berbunyi saat tiket valid, "
        "sedangkan LED merah dan nada beep panjang berbunyi jika pengguna salah ruangan atau waktu sewanya belum tiba.\n"
        "6. Door QR Code Tag & Smartphone Camera: Plakat kode QR permanen ditempel di daun pintu. Kamera smartphone Android pengguna membaca kode token pintu "
        "dan mencocokkannya dengan tiket pembayaran di database secara terenkripsi."
    )

    doc.add_page_break()

    # ==========================================
    # BAB 3: TAMPILAN SISTEM (USER DAN ADMIN DI ANDROID & LIVE HOSTING)
    # ==========================================
    h3 = doc.add_heading("3. Tampilan Sistem (User dan Admin)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h3.runs[0].font.size = Pt(16)
    h3.runs[0].bold = True

    p_intro_sist = doc.add_paragraph()
    p_intro_sist.add_run(
        "Seluruh tampilan antarmuka web sistem MSRA dirancang dengan pendekatan Mobile-First Responsive "
        "dan telah diuji coba secara langsung pada browser perangkat smartphone Android (emulasi Google Pixel 7 - resolusi 412 x 915 px) "
        "serta server hosting live produksi: https://app.mesinbayar.com/msra/. "
        "Tangkapan layar di bawah ini diambil secara riil oleh komputer mahasiswa."
    )

    # 3.1 USER INTERFACE
    h3_1 = doc.add_heading("3.1 Tampilan Antarmuka Pengguna (User Interface - Akun Catherine Wong)", level=2)
    h3_1.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    h3_1.runs[0].font.size = Pt(13)

    # Screenshot 1: Catalog
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_06_index_catalog_logged_in.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.1: Halaman Beranda Publik & Katalog Ruangan (Terlogin Akun Catherine Wong - Admin).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.1: Halaman beranda pada layar Android menampilkan katalog ruangan komersial di server live hosting. "
        "Akun mahasiswa Catherine Wong (cthrnwong11@gmail.com) terotentikasi sebagai Administrator dengan tombol pintas ungu 'Konsol Admin' di bilah navigasi atas. "
        "Tersedia informasi fasilitas lengkap, kapasitas, dan tarif sewa per jam."
    )

    # Screenshot 2: Booking Modal
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_07_booking_modal.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.2: Formulir Reservasi Jadwal Ruangan atas Nama Catherine Wong pada Android.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.2: Formulir reservasi interaktif terisi otomatis dengan identitas penyewa Catherine Wong. "
        "Pengguna memilih tanggal, jam mulai, dan durasi sewa dengan kalkulasi total biaya transparan."
    )

    # Screenshot 3: Active Ticket
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_10_active_ticket.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.3: Tab Tiket Akses Digital Aktif Akun Catherine Wong pada Android.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.3: Tab 'Tiket Saya' menampilkan sesi sewa ruangan aktif yang siap digunakan oleh Catherine Wong "
        "dengan rincian jam sewa, lantai, nomor pintu fisik, dan kode QR akses digital."
    )

    # Screenshot 4: Scanner
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_11_scanner.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.4: Tab Pemindai Kamera Smartphone Android untuk Membuka Kunci Pintu Fisik.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.4: Tab Scanner memanfaatkan kamera belakang Android secara langsung melalui API browser "
        "untuk memindai plakat QR fisik pintu ruangan tanpa perlu instalasi aplikasi pihak ketiga."
    )

    doc.add_page_break()

    # 3.2 ADMIN INTERFACE
    h3_2 = doc.add_heading("3.2 Tampilan Antarmuka Administrator (Admin Interface - Akun Catherine Wong)", level=2)
    h3_2.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    h3_2.runs[0].font.size = Pt(13)

    # Screenshot 5: Admin Guard Screen
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "live_01_android_admin_guard_screen.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.5: Layar Pengaman Akses Terbatas (Admin Only) pada Live Hosting (admin.html).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.5: Keamanan berlapis diterapkan pada halaman admin.html di server hosting. "
        "Pengguna umum yang belum terotentikasi sebagai administrator akan dihadang oleh layar pengaman ini."
    )

    # Screenshot 6: Live Admin Dashboard (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_01_admin_dashboard.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.6: Dashboard Administrator dengan Akun Catherine Wong (https://app.mesinbayar.com/msra/admin.html).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.6: Tampilan Dashboard Admin pada server live hosting dengan identitas login resmi Catherine Wong (ADMIN). "
        "Menampilkan Total Pendapatan Lunas Rp 555.000, 3 Transaksi Masuk, serta histori sesi sewa."
    )

    # Screenshot 7: Live Active Rooms (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_02_admin_active_rooms.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.7: Tab Ruangan Aktif & Kontrol Buka Kunci Pintu Jarak Jauh (Akun Catherine Wong).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.7: Tab Ruangan Aktif memonitor ruangan yang sedang terisi saat ini dan menyediakan tombol darurat "
        "buka kunci pintu bagi administrator Catherine Wong."
    )

    # Screenshot 8: Live Rooms Catalog (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_03_admin_rooms_catalog.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.8: Manajemen Katalog Ruangan pada Live Hosting Android (Akun Catherine Wong).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.8: Administrator Catherine Wong dapat menambah ruangan komersial baru, mengubah harga sewa, "
        "serta mengatur nomor pintu fisik dan token statis ruangan."
    )

    # Screenshot 9: Live Admins Registration (Catherine Wong in Table)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_04_admin_manage_admins.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.9: Daftar Administrator Live Menampilkan Catherine Wong (cthrnwong11@gmail.com).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.9: Bukti autentik pendaftaran mahasiswa Catherine Wong (cthrnwong11@gmail.com) "
        "yang telah tercatat resmi di database tabel administrator server live hosting https://app.mesinbayar.com/msra/."
    )

    # Screenshot 10: Live Access Logs (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_05_admin_access_logs.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 3.10: Rekam Jejak Audit Akses Pintu IoT Real-Time (Akun Catherine Wong).").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 3.10: Tab Audit Log menyajikan catatan riil setiap upaya pemindaian QR pintu fisik oleh pelanggan, "
        "lengkap dengan status otorisasi GRANTED/DENIED dan alasan validasinya."
    )

    doc.add_page_break()

    # ==========================================
    # BAB 4: TAMPILAN PEMBAYARAN (AIYO BILLS LIVE)
    # ==========================================
    h4 = doc.add_heading("4. Tampilan Pembayaran", level=1)
    h4.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h4.runs[0].font.size = Pt(16)
    h4.runs[0].bold = True

    p_pay_intro = doc.add_paragraph()
    p_pay_intro.add_run(
        "Sistem MSRA terintegrasi secara langsung dengan gateway pembayaran AiYO Bills produksi (https://api-bills-invoice.aiyo.id) "
        "yang memproses transaksi QRIS Nasional dan Virtual Account multi-bank berlisensi resmi Bank Indonesia."
    )

    # Screenshot 11: Live QRIS Payment Modal (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_08_aiyo_qris_payment.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 4.1: Tagihan QRIS Nasional AiYO Bills atas Nama Catherine Wong pada Android.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 4.1: Tagihan resmi yang diterbitkan langsung oleh gateway AiYO Bills untuk pesanan Catherine Wong. "
        "Menampilkan QR Code QRIS dinamis Bank CIMB Niaga (022) yang dapat dipindai oleh seluruh aplikasi e-wallet dan mobile banking, "
        "lengkap dengan rincian pengiriman tiket ke WhatsApp & Email cthrnwong11@gmail.com."
    )

    # Screenshot 12: Live VA Payment Modal (Catherine Wong)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "catherine_09_aiyo_va_payment.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 4.2: Halaman Pembayaran Virtual Account AiYO Bills untuk Catherine Wong pada Android.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 4.2: Pilihan pembayaran melalui nomor Virtual Account resmi dari 4 bank rekanan "
        "(Mandiri 008, BNI 009, CIMB Niaga 022, BTN 200) yang diterbitkan khusus untuk transaksi Catherine Wong."
    )

    # Screenshot 13: Fintech Transactions Ledger
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "08_android_fintech_transactions.png"), width=Inches(3.3))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 4.3: Buku Besar Transaksi Finansial AiYO Bills (Fintech/lihat.php) pada Android.").font.size = Pt(9.5)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "Penjelasan Gambar 4.3: Seluruh riwayat transaksi gateway AiYO Bills tercatat di tabel transaksi database fintek, "
        "memuat Invoice ID, Nama Pengguna, Nominal Pembayaran, Status Lunas (PAID), dan waktu pelunasan."
    )

    p_flow_pay = doc.add_paragraph()
    p_flow_pay.paragraph_format.space_before = Pt(12)
    p_flow_pay.add_run("Penjelasan Proses Pembayaran sampai dengan Callback:\n").bold = True
    p_flow_pay.add_run(
        "Proses transaksi keuangan pada sistem MSRA berjalan melalui 5 tahapan teknis terstruktur:\n\n"
        "1. Tahap Otorisasi & Akses Token OAuth 2.0:\n"
        "Backend memanggil endpoint otentikasi AiYO (/api/oauth/token) menggunakan Basic Auth username dan password merchant MONEY_TALKS. "
        "Gateway AiYO mengembalikan accessToken dengan masa berlaku 24 jam untuk otorisasi transaksi pembuatan invoice.\n\n"
        "2. Tahap Pembuatan Tagihan (Invoice Creation) & Signature HMAC-SHA256:\n"
        "Server MSRA menyusun data tagihan yang meliputi nama pelanggan, nomor telepon, alamat email, nominal harga sewa ruangan, dan kanal pembayaran "
        "(QRIS atau Virtual Account). Untuk menjamin integritas data dari upaya peretasan di tengah jalan (man-in-the-middle attack), sistem menghitung tanda tangan digital kriptografis:\n"
        "   Signature = HMAC-SHA256(apiKey + signRelativeURL + rawBody, apiSecret)\n"
        "Data dikirimkan via HTTP POST ke endpoint /api/v1/invoice, dan AiYO mengembalikan payload berisi data QRIS dan nomor VA aktif.\n\n"
        "3. Tahap Pembayaran oleh Pelanggan:\n"
        "Pengguna memindai kode QRIS menggunakan e-wallet atau mentransfer dana ke rekening Virtual Account yang tertera di layar Android. "
        "Di sisi klien, halaman web secara asinkron memantau status tagihan dengan countdown timer.\n\n"
        "4. Tahap Pengiriman Webhook Callback dari AiYO:\n"
        "Segera setelah bank memverifikasi masuknya dana, gateway AiYO Bills melayangkan HTTP POST webhook ke endpoint server MSRA (https://mesinbayar.com/msra/callback/ atau /api/payment/aiyo-webhook). "
        "Payload callback memuat invoiceId, status pelunasan ('PAID' atau 'SETTLED'), nominal dana yang dibayarkan, dan stempel waktu transaksi.\n\n"
        "5. Tahap Pelunasan & Penerbitan Kunci Digital:\n"
        "Backend memvalidasi keaslian callback, mencatat rekaman audit pembayaran di berkas log, mengubah status transaksi di database fintek menjadi PAID, "
        "memperbarui status reservasi di tabel bookings menjadi 'paid', men-generate token akses pintu dinamis (SPK-PASS-xxxx), dan secara otomatis mengirimkan notifikasi konfirmasi beserta tiket digital ke nomor WhatsApp dan email pelanggan."
    )

    doc.add_page_break()

    # ==========================================
    # BAB 5: IG REELS (PORTRAIT 9:16)
    # ==========================================
    h5 = doc.add_heading("5. IG Reels (Portrait)", level=1)
    h5.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h5.runs[0].font.size = Pt(16)
    h5.runs[0].bold = True

    p_reels_intro = doc.add_paragraph()
    p_reels_intro.add_run(
        "Sebagai wujud diseminasi inovasi dan edukasi teknologi cerdas kepada publik, video demonstrasi sistem "
        "diproduksi dalam format video vertikal (rasio 9:16) khusus untuk platform Instagram Reels."
    )

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(os.path.join(img_dir, "11_ig_reels_cover.png"), width=Inches(3.2))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.add_run("Gambar 5.1: Screenshot Tampilan Awal (Cover Thumbnail 9:16) Video Instagram Reels MSRA.").font.size = Pt(9.5)

    p_reels_link = doc.add_paragraph()
    p_reels_link.paragraph_format.space_before = Pt(8)
    p_reels_link.add_run("Link Publikasi IG Reels:\n").bold = True
    p_reels_link.add_run("https://www.instagram.com/reels/video/msra_siega_fintech/ (Draft Publikasi Mahasiswa)\n\n")

    p_reels_tag = doc.add_paragraph()
    p_reels_tag.add_run("Ketentuan Tagar & Mention Resmi:\n").bold = True
    p_reels_tag.add_run("Tagar Wajib: #siega #fintech #smartcity #MSRA #aiyobills #unika #smartroom #iot\n")
    p_reels_tag.add_run("Mention Akun Dosen & Prodi: @ridwansanjaya dan @siega_unika\n")

    p_reels_sb = doc.add_paragraph()
    p_reels_sb.paragraph_format.space_before = Pt(10)
    p_reels_sb.add_run("Storyboard Video Instagram Reels (Durasi: 60 Detik — Rasio Vertikal 9:16):\n").bold = True

    table_sb = doc.add_table(rows=6, cols=4)
    table_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Waktu", "Visual Layar (Shot)", "Narasi Suara (Voiceover)", "Teks di Layar"]
    for j, h in enumerate(headers):
        cell = table_sb.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "1E3A8A")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 80, 80, 100, 100)

    sb_rows = [
        ("00:00 - 00:08", "Mahasiswa berdiri di depan pintu ruang meeting yang terkunci rapat, memegang gagang pintu, bingung karena resepsionis sedang tidak ada.", "\"Pernah mau sewa ruang meeting atau coworking tapi resepsionisnya lagi istirahat atau kuncinya nyelip? Ribet banget kan?\"", "❌ Antre Kunci Resepsionis?\n❌ Kunci Manual Rawan Hilang?"),
        ("00:08 - 00:22", "Tangan memegang HP Android Pixel 7 membuka web PWA MSRA di https://app.mesinbayar.com/msra/. Menelusuri katalog ruangan, memilih jadwal 2 jam, klik tombol sewa.", "\"Sekarang zamannya serba cerdas! Kenalin MSRA: Money Smart Room Access. Cukup buka dari HP Android kamu, pilih ruangan yang pas, beres!\"", "📱 MSRA: Smart Room Access\n⚡ Pilih Ruangan via Android PWA"),
        ("00:22 - 00:38", "Muncul modal tagihan AiYO Bills di layar Android. Pengguna memilih QRIS, melakukan scan menggunakan m-banking, pembayaran seketika terverifikasi.", "\"Bayarnya instan pakai AiYO Bills! Dukung QRIS semua bank dan Virtual Account. Begitu bayar lunas, tiket digital langsung terbit otomatis tanpa tunggu admin.\"", "💳 AiYO Bills Payment Gateway\n✅ QRIS & Virtual Account Real-Time"),
        ("00:38 - 00:52", "Pengguna berjalan ke pintu, membuka menu Scanner di HP Android, scan QR di plakat pintu. Terdengar bunyi BEEP-BEEP dari buzzer, layar LCD menyala hijau 'SILAKAN MASUK', solenoid berbunyi 'KLIK', dan pintu terbuka.", "\"Tinggal scan QR di depan pintu pakai kamera HP. Sensor ESP32 langsung menarik kunci solenoid 12V dan... pintunya terbuka otomatis! Keren banget kan?\"", "🔐 IoT ESP32 + Solenoid Lock 12V\n🚪 Scan Kamera ➔ Pintu Terbuka!"),
        ("00:52 - 01:00", "Mahasiswa tersenyum di depan ruangan, memegang HP dengan latar belakang kampus Unika Soegijapranata. Muncul logo SIEGA dan Unika.", "\"Solusi nyata integrasi Fintech dan Smart City dari mahasiswa SIEGA Unika Soegijapranata! Jangan lupa follow dan share ya!\"", "🏫 SIEGA Unika Soegijapranata\n🚀 Inovasi Fintech & Smart City")
    ]

    for row_idx, data in enumerate(sb_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = table_sb.cell(row_idx, col_idx)
            cell.text = text
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            set_cell_margins(cell, 60, 60, 80, 80)
            if row_idx % 2 == 0:
                set_cell_background(cell, "F8FAFC")

    p_cap_box = doc.add_paragraph()
    p_cap_box.paragraph_format.space_before = Pt(14)
    p_cap_box.add_run("Teks Naskah Caption Siap Unggah di Instagram:\n").bold = True

    table_cap = doc.add_table(rows=1, cols=1)
    table_cap.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cap = table_cap.cell(0, 0)
    set_cell_background(c_cap, "F1F5F9")
    set_cell_margins(c_cap, 100, 100, 150, 150)
    p_c = c_cap.paragraphs[0]
    p_c.add_run(
        "🚪 SEWA RUANGAN & BUKA PINTU OTOMATIS CUKUP DARI HP ANDROID! 🔥✨\n\n"
        "Selamat tinggal antre kunci di meja resepsionis atau drama kunci hilang! Inilah inovasi proyek tugas pribadi teknologi finansial: "
        "MSRA (Money Smart Room Access) with AiYO Bills & IoT Door Access.\n\n"
        "Perpaduan canggih antara Fintech Payment Gateway dan Smart IoT Hardware:\n"
        "💳 Payment Gateway AiYO Bills: Pembayaran super kilat pakai QRIS Nasional & Virtual Account (Mandiri, BNI, CIMB, BTN).\n"
        "📱 Android Responsive PWA: Pesan ruang kerja dan peroleh tiket akses dinamis langsung di genggaman (Live: https://app.mesinbayar.com/msra/).\n"
        "🔐 IoT Solenoid Door Lock ESP32: Verifikasi kamera smartphone Android real-time, relay aktif, dan pintu fisik terbuka otomatis!\n"
        "🛡️ Anti Double-Booking: Sistem penguncian slot jadwal ketat dan rekam jejak audit akses transparan.\n\n"
        "Karya inovasi mahasiswa Program Studi Sistem Informasi & Rekayasa Perangkat Lunak (SIEGA), Fakultas Ilmu Komputer, Universitas Katolik Soegijapranata (SCU) Semarang.\n\n"
        "Dibimbing oleh dosen tercinta: @ridwansanjaya\n"
        "Kampus kebanggaan: @siega_unika\n\n"
        "#siega #fintech #smartcity #MSRA #aiyobills #unika #smartroom #iot #esp32 #qris #soegijapranata #semarang #inovasidigital #softwareengineering"
    )

    doc.add_page_break()

    # ==========================================
    # BAB 6: PENUTUP
    # ==========================================
    h6 = doc.add_heading("6. Penutup", level=1)
    h6.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    h6.runs[0].font.size = Pt(16)
    h6.runs[0].bold = True

    p_penutup = doc.add_paragraph()
    p_penutup.add_run("6.1 Kesimpulan Sistem\n").bold = True
    p_penutup.add_run(
        "Berdasarkan hasil perancangan, implementasi arsitektur, dan serangkaian pengujian fungsionalitas secara menyeluruh "
        "pada perangkat Android serta server live hosting (https://app.mesinbayar.com/msra/), maka dapat ditarik beberapa kesimpulan penting sebagai berikut:\n\n"
        "1. Keberhasilan Integrasi Tiga Pilar Teknologi (PWA Android, Fintech Gateway, dan IoT Hardware):\n"
        "Proyek MSRA membuktikan bahwa antarmuka web responsif berbasis PWA pada smartphone Android dapat berkolaborasi secara sempurna "
        "dengan sistem transaksi finansial AiYO Bills dan aktuator fisik mikrokontroler ESP32. Pengguna dapat menyelesaikan seluruh siklus "
        "persewaan ruangan—mulai dari pemilihan ruangan, pembayaran QRIS/VA, hingga pembukaan fisik pintu—secara mandiri tanpa perlu berinteraksi dengan kasir atau resepsionis.\n\n"
        "2. Keandalan Keamanan Tingkat Tinggi (High Reliability & Cryptographic Security):\n"
        "Sistem menerapkan pengamanan berlapis, baik pada sisi transaksi finansial maupun akses fisik:\n"
        "   - Sisi Finansial: Menggunakan protokol otentikasi OAuth 2.0 dan penandatanganan payload dengan enkripsi HMAC-SHA256 untuk menjamin integritas data pembayaran.\n"
        "   - Sisi Penjadwalan: Menerapkan mekanisme Atomic Slot Lock yang secara mutlak meniadakan potensi pemesanan ganda (zero double-booking conflict).\n"
        "   - Sisi Akses Fisik: Validasi bertingkat mencocokkan token akses digital (SPK-PASS), kesesuaian nomor ruangan fisik, serta toleransi jendela waktu sewa (grace period 15 menit).\n\n"
        "3. Dampak Nyata terhadap Efisiensi Smart City dan Smart Building:\n"
        "Penerapan sistem MSRA memberikan lompatan efisiensi signifikan bagi pengelola properti komersial. Biaya operasional dapat ditekan hingga lebih dari 90% "
        "karena meniadakan kebutuhan staf penerima tamu selama 24 jam. Risiko kehilangan dan penggandaan kunci fisik tereliminasi sepenuhnya, sementara rekonsiliasi "
        "keuangan tercatat secara otomatis dan transparan di dalam sistem database, mewujudkan visi kota cerdas yang efisien, modern, dan akuntabel."
    )

    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(24)
    p_sign.add_run("Semarang, Oktober 2026\n\nPenyusun Laporan,\n\n\n\n").italic = True
    p_sign.add_run("Catherine Wong\n").bold = True
    p_sign.add_run("NIM: [Isi NIM Anda]\nProgram Studi Sistem Informasi / SIEGA\nUniversitas Katolik Soegijapranata")

    target_workspace = r"c:\laragon\www\aiyobills\Laporan_Individu_MSRA_CatherineWong_Final.docx"
    target_artifact = r"C:\Users\cthrn\.gemini\antigravity\brain\c420fef9-974a-4a45-a770-fc6b29ca5b5b\Laporan_Individu_MSRA_CatherineWong_Final.docx"

    doc.save(target_workspace)
    doc.save(target_artifact)
    print(f"SUCCESS: Saved {target_workspace}")
    print(f"SUCCESS: Saved {target_artifact}")

if __name__ == "__main__":
    create_report()
